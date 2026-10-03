from copy import deepcopy
from pathlib import Path

import pytest

from jev_eval.backend import request_for
from jev_eval.excel import cases_from_tables, read_tables, scalar, state_from
from jev_eval.schema import load_cases

ROOT = Path(__file__).parents[1]


@pytest.mark.parametrize("excel,count", [
    ("data/evaluation.xlsx", 600), ("data/smoke.xlsx", 17),
])
def test_workbook_cases_load_without_label_leakage(excel, count):
    actual = load_cases(ROOT / excel)
    assert len(actual) == count
    assert all(set(request_for(c)) == {"state", "questions"} for c in actual)


def test_edits_are_used_and_gold_does_not_leak():
    tables = read_tables(ROOT / "data/evaluation.xlsx")
    edited = deepcopy(tables)
    edited["用例"][0]["用户请求"] = "只搜索专辑"
    edited["用例"][0]["预期结果"] = "album"
    actual = cases_from_tables(edited)[0]
    assert actual.state["query"] == "只搜索专辑"
    assert actual.gold == "album"
    assert "gold" not in request_for(actual)


def test_bad_option_and_unknown_input_case_are_rejected():
    tables = read_tables(ROOT / "data/smoke.xlsx")
    bad = deepcopy(tables)
    bad["用例"][0]["预期结果"] = "not_an_option"
    with pytest.raises(ValueError, match="available option"):
        cases_from_tables(bad)
    bad = deepcopy(tables)
    bad["输入字段"][0]["用例ID"] = "missing"
    with pytest.raises(ValueError, match="unknown case ID"):
        cases_from_tables(bad)


def test_field_types_empty_containers_and_duplicates():
    fields = [
        {"字段路径": "/a", "类型": "array", "值": None},
        {"字段路径": "/a/0", "类型": "boolean", "值": "假"},
        {"字段路径": "/b", "类型": "null", "值": None},
        {"字段路径": "/a~1b", "类型": "integer", "值": 0},
    ]
    assert state_from("q", fields) == {"query": "q", "a": [False], "b": None, "a/b": 0}
    with pytest.raises(ValueError, match="Duplicate"):
        state_from("q", fields + [fields[0]])
    with pytest.raises(ValueError, match="parent"):
        state_from("q", fields[1:])
    with pytest.raises(ValueError, match="integral"):
        scalar("integer", 0.5)
    with pytest.raises(ValueError, match="boolean"):
        scalar("boolean", "maybe")


def test_native_controls_are_present():
    from openpyxl import load_workbook

    book = load_workbook(ROOT / "data/evaluation.xlsx")
    try:
        assert book["用例"].freeze_panes == "A2"
        assert len(book["用例"].data_validations.dataValidation) == 3
        assert len(book["用例"].tables) == 1
        assert book["用例"].max_row == 601
    finally:
        book.close()
