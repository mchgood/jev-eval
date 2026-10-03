"""Read editable workbook tables; no JSON cells and no formula execution."""

import math
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

from .schema import Case

HEADERS = {
    "用例": ["用例ID", "领域", "任务", "能力", "划分", "用户请求", "判断问题",
           "预期结果", "标注依据", "排序组", "候选ID", "标签"],
    "输入字段": ["用例ID", "字段路径", "类型", "值"],
    "判断标准": ["用例ID", "选项或等级", "说明"],
}


def read_tables(path):
    book = load_workbook(path, read_only=True, data_only=False)
    tables = {}
    try:
        for name, headers in HEADERS.items():
            if name not in book.sheetnames:
                raise ValueError(f"Missing worksheet: {name}")
            sheet = book[name]
            records = []
            for cells in sheet.iter_rows():
                if any(c.data_type == "f" for c in cells):
                    raise ValueError(f"{name}:{cells[0].row}: formulas are not allowed")
                values = [c.value for c in cells]
                if not any(v is not None for v in values):
                    continue
                if not records and values[:len(headers)] != headers:
                    raise ValueError(f"{name}: invalid headers; expected {headers}")
                records.append(values[:len(headers)])
            tables[name] = [dict(zip(headers, row)) for row in records[1:]]
    finally:
        book.close()
    return tables


def scalar(kind, value):
    if kind == "object":
        if value is not None:
            raise ValueError("object value must be blank")
        return {}
    if kind == "array":
        if value is not None:
            raise ValueError("array value must be blank")
        return []
    if kind == "null":
        if value is not None:
            raise ValueError("null value must be blank")
        return None
    if kind == "text":
        return "" if value is None else str(value)
    if kind == "boolean":
        if type(value) is bool:
            return value
        if value in ("true", "真"):
            return True
        if value in ("false", "假"):
            return False
        raise ValueError("boolean must be true/false or 真/假")
    if kind in ("integer", "number"):
        if type(value) not in (int, float) or not math.isfinite(value):
            raise ValueError("numeric value must be a finite Excel number")
        if kind == "integer" and value != int(value):
            raise ValueError("integer value must be integral")
        return int(value) if kind == "integer" else float(value)
    raise ValueError(f"Unknown field type: {kind}")


def state_from(query, fields):
    """Parent container rows make object keys and array indices unambiguous."""
    nodes = {"": {"query": query}}
    paths = set()
    for row in sorted(fields, key=lambda r: str(r["字段路径"]).count("/")):
        path = row["字段路径"]
        if not isinstance(path, str) or not path.startswith("/") or path == "/query":
            raise ValueError("Invalid field path; edit query in 用例 instead")
        if path in paths:
            raise ValueError(f"Duplicate field path: {path}")
        paths.add(path)
        parent, _, key = path.rpartition("/")
        key = key.replace("~1", "/").replace("~0", "~")
        if parent not in nodes:
            raise ValueError(f"Missing parent container: {parent}")
        container = nodes[parent]
        value = scalar(row["类型"], row["值"])
        if isinstance(container, dict):
            container[key] = value
        elif isinstance(container, list):
            if not key.isdigit() or str(int(key)) != key or int(key) != len(container):
                raise ValueError(f"Array indices must be consecutive from 0: {path}")
            container.append(value)
        else:
            raise TypeError(f"Parent is not a container: {parent}")
        nodes[path] = value
    return nodes[""]


def cases_from_tables(tables):
    inputs, criteria = defaultdict(list), defaultdict(list)
    ids = {r["用例ID"] for r in tables["用例"]}
    for name, target in (("输入字段", inputs), ("判断标准", criteria)):
        for row in tables[name]:
            if row["用例ID"] not in ids:
                raise ValueError(f"{name}: unknown case ID {row['用例ID']}")
            target[row["用例ID"]].append(row)
    cases = []
    for row in tables["用例"]:
        ident, kind = row["用例ID"], row["能力"]
        try:
            question = {"type": kind, "instructions": row["判断问题"]}
            gold = row["预期结果"]
            rules = criteria[ident]
            if kind == "choice":
                options = [r["选项或等级"] for r in rules]
                if any(not isinstance(o, str) or not o for o in options):
                    raise ValueError("Choice option must be nonempty text")
                if len(options) != len(set(options)):
                    raise ValueError("Duplicate Choice option")
                question["criteria"] = {r["选项或等级"]: r["说明"] for r in rules}
            elif kind == "score":
                levels = [r["选项或等级"] for r in rules]
                if any(type(n) not in (int, float) or n != int(n) for n in levels):
                    raise ValueError("Score levels must be integers")
                if sorted(levels) != list(range(len(levels))):
                    raise ValueError("Score levels must be consecutive from 0")
                question["criteria"] = [r["说明"] for r in sorted(rules, key=lambda r: r["选项或等级"])]
                gold = scalar("integer", gold)
            elif kind == "noul":
                if rules:
                    if {r["选项或等级"] for r in rules} != {"true", "false"} or len(rules) != 2:
                        raise ValueError("Noul criteria must contain true and false once")
                    question["criteria"] = {r["选项或等级"]: r["说明"] for r in rules}
                gold = None if gold in (None, "未知") else scalar("boolean", gold)
            cases.append(Case.model_validate({
                "id": ident, "domain": row["领域"], "task": row["任务"],
                "split": row["划分"], "state": state_from(row["用户请求"], inputs[ident]),
                "question": question, "gold": gold, "rationale": row["标注依据"],
                "ranking_group": row["排序组"] or None, "candidate_id": row["候选ID"] or None,
                "tags": str(row["标签"] or "").split(";") if row["标签"] else [],
            }))
        except Exception as exc:
            raise ValueError(f"用例 {ident}: {exc}") from exc
    return cases


def load_excel(path: Path):
    return cases_from_tables(read_tables(path))
